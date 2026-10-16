import pytest

from course_kb import cli
from course_kb.io_utils import read_json


@pytest.fixture
def run_cli(ctx, project, monkeypatch):
    monkeypatch.setattr(cli, "build_context", lambda ws: ctx)
    monkeypatch.chdir(project)

    def run(*args):
        return cli.main(["--workspace", str(project / "workspace"), *args])
    return run


def test_cli_full_flow(run_cli, ctx, project):
    assert run_cli("run", str(project / "sources.yaml"), "--audience", "大學二年級，已修過微積分",
                   "--title", "機器學習導論：最佳化") == 0
    assert run_cli("slides") == 2  # not approved yet
    n = ctx.fake_llm.count()
    assert run_cli("brief", "--render-only") == 0
    assert run_cli("brief", "--approve") == 0
    assert ctx.fake_llm.count() == n
    assert run_cli("slides") == 0
    assert read_json(ctx.ws.slides_json)["status"] == "complete"
    assert read_json(ctx.ws.brief_json)["audience"] == "大學二年級，已修過微積分"


def test_cli_render_and_approve_are_exclusive(run_cli):
    with pytest.raises(SystemExit):
        run_cli("brief", "--render-only", "--approve")


def test_cli_partial_exit_code(run_cli, ctx, project):
    ctx.fake_llm.fail_chunks.add("notes-a-001")
    assert run_cli("run", str(project / "sources.yaml")) == 1
    assert run_cli("brief", "--approve") == 2


def test_cli_rejects_duplicate_source_ids(run_cli, project):
    (project / "dup.yaml").write_text(
        "sources:\n  - {id: a, type: text, location: materials/a.md}\n"
        "  - {id: a, type: text, location: materials/b.md}\n", encoding="utf-8")
    assert run_cli("ingest", str(project / "dup.yaml")) == 2


def test_youtube_requires_rights_confirmed(run_cli, ctx, project):
    (project / "yt.yaml").write_text(
        "sources:\n  - {id: yt, type: youtube, location: 'https://youtu.be/x'}\n", encoding="utf-8")
    assert run_cli("ingest", str(project / "yt.yaml")) == 1
    units = read_json(project / "workspace" / "reports" / "run.json")["stages"]["ingest"]["units"]
    assert units["yt"]["reason"] == "PermissionError"
    assert ctx.fake_tools.calls == []  # nothing is queried or downloaded without confirmation


def _deck_text(path):
    from pptx import Presentation
    prs = Presentation(str(path))
    return prs, ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]


def test_cli_build_one_command(run_cli, ctx, project):
    src = project / "sources.yaml"
    src.write_text("course:\n  title: 最佳化入門\n  audience: 大學生\n  objectives: 說明梯度下降\n"
                   + src.read_text(encoding="utf-8"), encoding="utf-8")
    assert run_cli("build", str(src)) == 0
    brief = read_json(ctx.ws.brief_json)
    assert (brief["title"], brief["audience"], brief["objectives"]) == ("最佳化入門", "大學生", ["說明梯度下降"])
    assert read_json(ctx.ws.slides_json)["status"] == "complete"
    prs, texts = _deck_text(project / "slides.pptx")
    assert "最佳化入門" in texts[0]
    assert "品質報告：complete" in texts[1]
    items = read_json(ctx.ws.slides_json)["slides"]
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    assert all(any(it["speaker_script"] in n for n in notes) for it in items)


def test_cli_build_partial_still_exports_and_flags_gaps(run_cli, ctx, project):
    ctx.fake_llm.fail_chunks.add("notes-a-001")
    assert run_cli("build", str(project / "sources.yaml")) == 1
    decisions = read_json(ctx.ws.brief_json)["review_decisions"]
    assert any(d["decision"] == "accept_missing" and "未經人工審核" in d["reason"] for d in decisions)
    _, texts = _deck_text(project / "slides.pptx")
    assert "PARTIAL" in texts[1]
    assert "上游缺漏" in texts[1] and "自動決定" in texts[1]


def test_cli_build_stop_before_approve(run_cli, ctx, project):
    assert run_cli("build", str(project / "sources.yaml"), "--stop-before-approve") == 0
    assert not ctx.ws.approval_file.exists()
    assert not (project / "slides.pptx").exists()
