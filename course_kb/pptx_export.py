"""Export slides.json to a .pptx deck. Partial results are exported with every gap made visible."""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from .brief import AUTO_REASON, BriefError
from .io_utils import read_json
from .state import Context

log = logging.getLogger("course_kb")

WARN = RGBColor(0xC0, 0x00, 0x00)
MUTED = RGBColor(0x59, 0x59, 0x59)
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.5)
REPORT_LINES_PER_SLIDE = 14


def _text(slide, left, top, width, height, lines: list[str], size: int, color=None, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.bold = bold
        if color is not None:
            p.font.color.rgb = color
        p.space_after = Pt(size // 3)
    return box


def _title(slide, text: str, color=None):
    _text(slide, MARGIN, Inches(0.3), W - 2 * MARGIN, Inches(1.0), [text], 30, color, bold=True)


def _body_size(lines: list[str]) -> int:
    n = sum(len(s) for s in lines)
    return 20 if n <= 250 else 16 if n <= 450 else 13


def _picture(slide, path: Path, left, top, max_w, max_h) -> None:
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    pw, ph = int(w * scale), int(h * scale)
    slide.shapes.add_picture(str(path), left + (max_w - pw) // 2, top, pw, ph)


def _quality_lines(doc: dict, outline: dict, brief: dict, cards: dict, slide_units: dict) -> list[str]:
    titles = {s["slide_id"]: s["title"] for s in outline.get("slides", [])}
    chapters = {c["chapter_id"]: c["title"] for c in brief.get("chapters", [])}
    L = []
    for e in doc["validation_errors"]:
        L.append(f"⚠ 驗證錯誤：{e}")
    for sid in doc["failed_slides"]:
        err = (slide_units.get(sid) or {}).get("error") or ""
        L.append(f"⚠ 生成失敗：{sid} {titles.get(sid, '')}{'（' + err[:80] + '）' if err else ''}")
    for cid in doc["failed_outline_chapters"]:
        L.append(f"⚠ 大綱失敗章節：{cid} {chapters.get(cid, '')}")
    for cid in doc["required_uncovered_card_ids"]:
        L.append(f"⚠ 未涵蓋的必要卡片：{cid} {cards.get(cid, {}).get('concept', '')}")
    for m in doc["missing"]:
        L.append(f"⚠ 上游缺漏：{m['stage']} / {m['unit']}（{m['reason']}）")
    for d in brief.get("review_decisions", []):
        if str(d.get("reason", "")).startswith(AUTO_REASON):
            L.append(f"△ 自動決定 {d.get('decision')}：{d['reason'][len(AUTO_REASON) + 1:]}")
    for it in doc["slides"]:
        if it["review_items"]:
            L.append(f"△ 待確認 {len(it['review_items'])} 項：{it['slide_id']} {it['title']}")
    return L


def export_pptx(ctx: Context, out: Optional[Path] = None) -> Path:
    ws = ctx.ws
    doc = read_json(ws.slides_json)
    outline = read_json(ws.outline_file)
    brief = read_json(ws.brief_json)
    if not doc or not outline or not brief:
        raise BriefError("no slides/slides.json; run 'slides' first")
    cards = {c["card_id"]: c for c in (read_json(ws.cards_file) or {}).get("cards", [])}
    slide_units = ctx.report().data["stages"].get("slides", {}).get("units", {})
    items = {it["slide_id"]: it for it in doc["slides"]}
    chapters = {c["chapter_id"]: c for c in brief.get("chapters", [])}

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    _text(s, MARGIN, Inches(2.2), W - 2 * MARGIN, Inches(1.5), [brief.get("title", "")], 40, bold=True)
    sub = [f"受眾：{brief.get('audience') or '（未指定）'}"] + [f"目標：{o}" for o in brief.get("objectives", [])]
    _text(s, MARGIN, Inches(3.8), W - 2 * MARGIN, Inches(2.5), sub, 20, MUTED)

    complete = doc["status"] == "complete"
    report = _quality_lines(doc, outline, brief, cards, slide_units)
    pages = [report[i:i + REPORT_LINES_PER_SLIDE] for i in range(0, len(report), REPORT_LINES_PER_SLIDE)] or [[]]
    for n, page in enumerate(pages, 1):
        s = prs.slides.add_slide(blank)
        head = "品質報告：complete" if complete else "⚠ 品質報告：PARTIAL — 內容有缺漏，請逐項確認"
        _title(s, head + (f"（{n}/{len(pages)}）" if len(pages) > 1 else ""), None if complete else WARN)
        _text(s, MARGIN, Inches(1.4), W - 2 * MARGIN, H - Inches(1.8), page or ["沒有待處理項目。"], 15)

    current = None
    for sl in outline["slides"]:
        if sl["chapter_id"] != current:
            current = sl["chapter_id"]
            s = prs.slides.add_slide(blank)
            _text(s, MARGIN, Inches(2.6), W - 2 * MARGIN, Inches(1.5),
                  [f"{current}  {chapters.get(current, {}).get('title', '')}"], 36, bold=True)
            objs = chapters.get(current, {}).get("learning_objectives", [])
            _text(s, MARGIN, Inches(4.0), W - 2 * MARGIN, Inches(2.5), objs, 18, MUTED)
        it = items.get(sl["slide_id"])
        s = prs.slides.add_slide(blank)
        if it is None:
            _title(s, f"⚠ 生成失敗：{sl['title']}", WARN)
            err = (slide_units.get(sl["slide_id"]) or {}).get("error") or "未知原因"
            _text(s, MARGIN, Inches(1.5), W - 2 * MARGIN, Inches(4),
                  [f"投影片 {sl['slide_id']} 未產生，請重跑 slides 補做。", f"學習目標：{sl['learning_objective']}",
                   f"卡片：{', '.join(sl['card_ids'])}", f"原因：{err}"], 18, WARN)
            s.notes_slide.notes_text_frame.text = f"生成失敗：{err}"
            continue
        flagged = bool(it["review_items"])
        _title(s, ("△ " if flagged else "") + it["title"], WARN if flagged else None)
        frame = next((v["frame"] for v in it["visuals"]
                      if v.get("frame") and (ws.root / v["frame"]["path"]).exists()), None)
        text_w = Inches(7.3) if frame else W - 2 * MARGIN
        _text(s, MARGIN, Inches(1.4), text_w, Inches(4.6), [f"• {c}" for c in it["content"]],
              _body_size(it["content"]))
        if frame:
            _picture(s, ws.root / frame["path"], MARGIN + Inches(7.5), Inches(1.5), Inches(4.8), Inches(4.4))
        hints = [v["description"] for v in it["visuals"] if not v.get("frame")]
        if hints:
            _text(s, MARGIN, Inches(6.1), W - 2 * MARGIN, Inches(1.2),
                  ["視覺建議：" + "；".join(hints)], 11, MUTED)
        notes = [it["speaker_script"], "", f"學習目標：{it['learning_objective']}",
                 f"卡片：{', '.join(it['card_ids'])}", "來源："]
        notes += [f"- {src['source_id']} {src.get('location') or ''} {src['locator']}" for src in it["sources"]]
        if flagged:
            notes += ["", "待確認："] + [f"- [{r['type']}] {r['text']}" for r in it["review_items"]]
        s.notes_slide.notes_text_frame.text = "\n".join(notes)

    out = Path(out) if out else Path.cwd() / "slides.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    log.info("pptx: %s (%s, %d slides)", out, doc["status"], len(prs.slides))
    return out
